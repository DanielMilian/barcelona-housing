import sys

from openpyxl import load_workbook

wb = load_workbook(sys.argv[1], read_only=True, data_only=True)
for ws in wb.worksheets:
    print(f"\n== {ws.title}")
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        print(row[:25])
        if i >= 3:
            break

wb.close()

